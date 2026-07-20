from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "GraphRAG与SAADS-Red-team项目技术说明.md"
OUTPUT = ROOT / "docs" / "GraphRAG与SAADS-Red-team项目技术说明_最终版.docx"
ASSETS = ROOT / "docs" / "assets"

PAGE_WIDTH_DXA = 12240
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "586878"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
INK = "20252B"
GRID = "C9D3DE"
ACCENT = "2D8C88"
GOLD = "B07A18"
RISK = "9B1C1C"


def find_font(bold: bool = False) -> Path:
    candidates = (
        [Path(r"C:\Windows\Fonts\msyhbd.ttc"), Path(r"C:\Windows\Fonts\simhei.ttf")]
        if bold
        else [Path(r"C:\Windows\Fonts\msyh.ttc"), Path(r"C:\Windows\Fonts\simsun.ttc")]
    )
    candidates += [Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf")]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("No usable CJK font found in C:\\Windows\\Fonts")


REGULAR_FONT = find_font(False)
BOLD_FONT = find_font(True)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(BOLD_FONT if bold else REGULAR_FONT), size=size)


def rounded_box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], fill: str, outline: str) -> None:
    draw.rounded_rectangle(box, radius=24, fill=f"#{fill}", outline=f"#{outline}", width=3)


def centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    size: int,
    color: str = INK,
    bold: bool = False,
) -> None:
    text_font = font(size, bold)
    bbox = draw.multiline_textbbox((0, 0), text, font=text_font, spacing=8, align="center")
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    x = box[0] + (box[2] - box[0] - width) / 2
    y = box[1] + (box[3] - box[1] - height) / 2 - 2
    draw.multiline_text((x, y), text, font=text_font, fill=f"#{color}", spacing=8, align="center")


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = BLUE) -> None:
    draw.line([start, end], fill=f"#{color}", width=6)
    x, y = end
    draw.polygon([(x, y), (x - 16, y - 10), (x - 16, y + 10)], fill=f"#{color}")


def make_architecture_figure() -> Path:
    path = ASSETS / "graphrag_lifecycle.png"
    canvas = Image.new("RGB", (2000, 920), f"#{WHITE}")
    draw = ImageDraw.Draw(canvas)
    draw.text((90, 54), "GraphRAG 的两阶段生命周期", font=font(48, True), fill=f"#{NAVY}")
    draw.text((92, 118), "离线构建结构化记忆，查询时按问题粒度选择上下文", font=font(25), fill=f"#{MUTED}")

    x_values = [80, 390, 700, 1010, 1320, 1630]
    labels = [
        "可信文档\n与元数据",
        "分块\nText units",
        "实体 / 关系\n抽取与归并",
        "社区发现\n与报告",
        "向量化\n与持久化",
        "可查询的\nGraphRAG 索引",
    ]
    fills = ["F4F6F9", "E8EEF5", "E6F2F1", "FFF3D8", "EDE8F5", "DDEAF7"]
    boxes = []
    for x, label, fill in zip(x_values, labels, fills):
        box = (x, 245, x + 250, 420)
        boxes.append(box)
        rounded_box(draw, box, fill, GRID)
        centered_text(draw, box, label, 28, NAVY, True)
    for left, right in zip(boxes, boxes[1:]):
        arrow(draw, (left[2] + 10, 332), (right[0] - 12, 332))

    draw.rounded_rectangle((80, 530, 1920, 805), radius=28, fill="#F8FAFC", outline=f"#{GRID}", width=3)
    draw.text((115, 565), "查询引擎", font=font(30, True), fill=f"#{NAVY}")
    modes = [
        ("BASIC", "文本块事实定位", BLUE),
        ("LOCAL", "实体邻域与局部关系", ACCENT),
        ("GLOBAL", "社区报告 map-reduce", GOLD),
        ("DRIFT", "广域起点 + 迭代多跳", "6E56A0"),
    ]
    x = 115
    for title, detail, color in modes:
        box = (x, 635, x + 405, 755)
        rounded_box(draw, box, WHITE, color)
        draw.text((x + 24, 657), title, font=font(28, True), fill=f"#{color}")
        draw.text((x + 24, 704), detail, font=font(21), fill=f"#{INK}")
        x += 445
    canvas.save(path, quality=95)
    return path


def make_project_figure() -> Path:
    path = ASSETS / "project_architecture.png"
    canvas = Image.new("RGB", (2000, 1080), f"#{WHITE}")
    draw = ImageDraw.Draw(canvas)
    draw.text((90, 52), "SAADS-Red-team 项目数据与控制流", font=font(48, True), fill=f"#{NAVY}")
    draw.text((92, 116), "密钥只在本地配置层加载；看板读取静态快照，不携带 API 密钥", font=font(25), fill=f"#{MUTED}")

    rows = [
        (210, "01 可信来源", ["MITRE ATLAS", "OWASP LLM Top 10", "PyRIT / garak"], "F4F6F9", BLUE),
        (370, "02 确定性采集", ["raw 快照", "Markdown + front matter", "manifest + SHA-256"], "E8EEF5", BLUE),
        (530, "03 GraphRAG Standard", ["分块与抽取", "社区与报告", "Parquet + LanceDB"], "E6F2F1", ACCENT),
        (690, "04 质量与评测", ["图谱分析", "40 题路由集", "四模式冒烟"], "FFF3D8", GOLD),
        (850, "05 使用界面", ["CLI 查询", "静态证据探索", "质量态势看板"], "EDE8F5", "6E56A0"),
    ]
    for y, label, items, fill, outline in rows:
        label_box = (90, y, 410, y + 112)
        rounded_box(draw, label_box, outline, outline)
        centered_text(draw, label_box, label, 28, WHITE, True)
        x = 475
        for item in items:
            box = (x, y, x + 430, y + 112)
            rounded_box(draw, box, fill, outline)
            centered_text(draw, box, item, 25, NAVY, True)
            x += 470
    for y1, y2 in zip([322, 482, 642, 802], [370, 530, 690, 850]):
        draw.line([(250, y1), (250, y2 - 16)], fill=f"#{MUTED}", width=5)
        draw.polygon([(250, y2), (238, y2 - 20), (262, y2 - 20)], fill=f"#{MUTED}")
    canvas.save(path, quality=95)
    return path


def make_bar_chart(
    filename: str,
    title: str,
    values: list[tuple[str, int]],
    total: int,
    colors: list[str],
    width: int = 1800,
) -> Path:
    height = 250 + len(values) * 115
    path = ASSETS / filename
    canvas = Image.new("RGB", (width, height), f"#{WHITE}")
    draw = ImageDraw.Draw(canvas)
    draw.text((70, 42), title, font=font(42, True), fill=f"#{NAVY}")
    max_value = max(value for _, value in values)
    label_x, bar_x, bar_width = 70, 480, width - 860
    for index, ((label, value), color) in enumerate(zip(values, colors)):
        y = 155 + index * 112
        draw.text((label_x, y + 5), label, font=font(25, True), fill=f"#{INK}")
        draw.rounded_rectangle((bar_x, y, bar_x + bar_width, y + 48), radius=16, fill="#EEF2F6")
        fill_width = max(10, int(bar_width * value / max_value))
        draw.rounded_rectangle((bar_x, y, bar_x + fill_width, y + 48), radius=16, fill=f"#{color}")
        value_text = f"{value:,}  ·  {value / total:.2%}"
        draw.text((bar_x + bar_width + 20, y + 4), value_text, font=font(23), fill=f"#{MUTED}")
    canvas.save(path, quality=95)
    return path


def generate_figures() -> dict[str, tuple[Path, str]]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    return {
        "architecture": (make_architecture_figure(), "图 1  GraphRAG 离线索引与四类查询模式"),
        "project": (make_project_figure(), "图 2  当前项目的端到端数据与控制流"),
        "sources": (
            make_bar_chart(
                "corpus_sources.png",
                "当前语料来源构成（297 篇）",
                [("MITRE ATLAS", 271), ("NVIDIA garak", 15), ("OWASP 2025", 10), ("Microsoft PyRIT", 1)],
                297,
                [BLUE, ACCENT, GOLD, "6E56A0"],
            ),
            "图 3  当前语料来源数量与占比",
        ),
        "entities": (
            make_bar_chart(
                "entity_types.png",
                "七类领域实体分布（不含 3 个越界实体）",
                [
                    ("攻击技术", 987),
                    ("组件", 949),
                    ("防御控制", 502),
                    ("弱点", 292),
                    ("工具", 114),
                    ("评估", 75),
                    ("标准", 51),
                ],
                2973,
                [RISK, BLUE, ACCENT, GOLD, "6E56A0", "4E7793", "707A84"],
            ),
            "图 4  当前七类领域实体的数量分布",
        ),
    }


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None, name: str = "Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size: float, color: str = INK, bold: bool = False, latin: str = "Calibri") -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def shade_paragraph(paragraph, fill: str, border: str | None = None) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border:
        pbdr = ppr.find(qn("w:pBdr"))
        if pbdr is None:
            pbdr = OxmlElement("w:pBdr")
            ppr.append(pbdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border)
        pbdr.append(left)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def set_table_borders(table, color: str = GRID, size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[int], indent: int = TABLE_INDENT_DXA) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell, **CELL_MARGINS_DXA)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table_widths(rows: list[list[str]]) -> list[int]:
    columns = len(rows[0])
    scores: list[float] = []
    for index in range(columns):
        values = [re.sub(r"[`*]", "", row[index]) for row in rows]
        longest = min(max(len(value) for value in values), 45)
        numeric = all(re.fullmatch(r"[\d,./%\s:+-]+", value.strip()) for value in values[1:] if value.strip())
        scores.append(7 if numeric else max(9, longest))
    total = sum(scores)
    widths = [max(900, int(CONTENT_WIDTH_DXA * score / total)) for score in scores]
    while sum(widths) > CONTENT_WIDTH_DXA:
        index = max(range(columns), key=lambda i: widths[i])
        widths[index] -= 1
    while sum(widths) < CONTENT_WIDTH_DXA:
        index = max(range(columns), key=lambda i: scores[i])
        widths[index] += 1
    return widths


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.extend([rfonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([rpr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*.+?\*\*|`[^`]+`|https?://[^\s)]+)")


def add_inline(paragraph, text: str, default_size: float | None = None, default_color: str = INK) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, default_size, color=default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, default_size, bold=True, color=default_color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, default_size or 9.5, color=DARK_BLUE, name="Consolas")
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        elif token.startswith("["):
            link_text, url = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, link_text, url)
        else:
            add_hyperlink(paragraph, token, token)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, default_size, color=default_color)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, 9, color=MUTED)


def add_numbering_definition(document: Document, kind: str) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"A{abstract_id:07X}")
    template = OxmlElement("w:tmpl")
    template.set(qn("w:val"), f"B{abstract_id:07X}")
    abstract.extend([nsid, template])
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.extend([tabs, indent, spacing])
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.append(rfonts)
    level.extend([start, num_fmt, lvl_text, suffix, justification, ppr, rpr])
    abstract.append(level)
    # OOXML requires every abstractNum to precede every num instance. Appending
    # an abstract definition after existing num elements makes Word repair the
    # numbering part and can merge bullets into a continuous decimal list.
    first_num_index = next(
        (index for index, child in enumerate(numbering) if child.tag == qn("w:num")),
        len(numbering),
    )
    numbering.insert(first_num_index, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_element])


def setup_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    set_style_font(normal, 11, INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    set_style_font(title, 30, NAVY, True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.0
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, 15, "2B5163")
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.paragraph_format.line_spacing = 1.15

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        set_style_font(style, size, color, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name, size, color, bold, alignment in (
        ("Cover Kicker", 10.5, GOLD, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("Cover Meta", 10, MUTED, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Figure Caption", 9, MUTED, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Code Block", 9, DARK_BLUE, False, WD_ALIGN_PARAGRAPH.LEFT),
        ("Table Text", 9.5, INK, False, WD_ALIGN_PARAGRAPH.LEFT),
    ):
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        set_style_font(style, size, color, bold, "Consolas" if name == "Code Block" else "Calibri")
        style.paragraph_format.alignment = alignment
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4 if name != "Code Block" else 0)
        style.paragraph_format.line_spacing = 1.1


def set_page_geometry(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def configure_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("GraphRAG × LLM 应用安全防御")
    set_run_font(run, 8.5, bold=True, color=MUTED)
    tab_stops = paragraph.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(6.5), alignment=2)
    run = paragraph.add_run("\t技术说明 · 2026-07-20")
    set_run_font(run, 8.5, color=MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run("第 ")
    set_run_font(run, 9, color=MUTED)
    add_field(paragraph, "PAGE")
    run = paragraph.add_run(" 页")
    set_run_font(run, 9, color=MUTED)


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(68)
    kicker = document.add_paragraph(style="Cover Kicker")
    kicker.add_run("TECHNICAL REFERENCE  ·  VERSION 1.0")
    kicker.paragraph_format.space_after = Pt(18)

    title = document.add_paragraph(style="Title")
    title.add_run("GraphRAG 与 SAADS-Red-team")
    title2 = document.add_paragraph(style="Title")
    title2.add_run("项目技术说明")
    title2.paragraph_format.space_after = Pt(12)

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("面向大模型应用安全防御的可审计知识图谱检索增强生成工程")
    subtitle.paragraph_format.space_after = Pt(42)

    for line in (
        "项目包：llm-defense-graphrag 0.1.0",
        "技术基线：Microsoft GraphRAG 3.1.1 · Python 3.12 · Standard Index",
        "项目快照：2026-07-20",
    ):
        paragraph = document.add_paragraph(style="Cover Meta")
        paragraph.add_run(line)

    callout = document.add_paragraph()
    callout.alignment = WD_ALIGN_PARAGRAPH.CENTER
    callout.paragraph_format.space_before = Pt(54)
    callout.paragraph_format.space_after = Pt(0)
    callout.paragraph_format.left_indent = Inches(0.55)
    callout.paragraph_format.right_indent = Inches(0.55)
    callout.paragraph_format.line_spacing = 1.15
    shade_paragraph(callout, LIGHT_BLUE, BLUE)
    run = callout.add_run("当前状态：Standard 索引已完成 · 四模式技术冒烟通过 · 完整 40 题人工评测待完成")
    set_run_font(run, 10.5, bold=True, color=NAVY)
    document.add_page_break()


def add_figure(document: Document, path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(path), width=Inches(6.35))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", caption)
    doc_pr.set("name", path.stem)
    caption_paragraph = document.add_paragraph(style="Figure Caption")
    caption_paragraph.add_run(caption)
    caption_paragraph.paragraph_format.space_after = Pt(8)


def add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    widths = table_widths(rows)
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        prevent_row_split(row)
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Table Text"]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if row_index == 0 or re.fullmatch(r"[\d,./%\s:+-]+", value.replace("**", "").strip())
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline(paragraph, value, 9.3, NAVY if row_index == 0 else INK)
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(2)


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    paragraph.paragraph_format.left_indent = Inches(0.16)
    paragraph.paragraph_format.right_indent = Inches(0.10)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.keep_together = True
    shade_paragraph(paragraph, LIGHT_GRAY, GRID)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run, 8.7, color=DARK_BLUE, name="Consolas")
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        if index < len(lines) - 1:
            run.add_break()


def add_callout(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.16)
    paragraph.paragraph_format.right_indent = Inches(0.10)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.2
    shade_paragraph(paragraph, CALLOUT, BLUE)
    add_inline(paragraph, text, 10.4, NAVY)


def parse_markdown(document: Document, text: str, figures: dict[str, tuple[Path, str]]) -> None:
    first_break = text.find("<!-- PAGEBREAK -->")
    body = text[first_break + len("<!-- PAGEBREAK -->") :] if first_break >= 0 else text
    lines = body.splitlines()
    bullet_num_id = add_numbering_definition(document, "bullet")
    ordered_num_id: int | None = None
    in_ordered = False
    index = 0

    def add_body_paragraph(content: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.25
        add_inline(paragraph, content.strip(), 11, INK)

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            in_ordered = False
            index += 1
            continue
        if stripped == "<!-- PAGEBREAK -->":
            document.add_page_break()
            in_ordered = False
            index += 1
            continue
        figure_match = re.fullmatch(r"<!-- FIGURE:([a-z]+) -->", stripped)
        if figure_match:
            path, caption = figures[figure_match.group(1)]
            add_figure(document, path, caption)
            index += 1
            continue
        if stripped.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows: list[list[str]] = []
            for row_index, line in enumerate(table_lines):
                cells = [cell.strip() for cell in line.strip("|").split("|")]
                if row_index == 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    continue
                rows.append(cells)
            if rows:
                add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_paragraph(style=f"Heading {level}")
            add_inline(paragraph, heading.group(2), None, BLUE if level < 3 else DARK_BLUE)
            index += 1
            continue
        if stripped.startswith(">"):
            text_parts: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                text_parts.append(lines[index].strip()[1:].strip())
                index += 1
            add_callout(document, " ".join(text_parts))
            continue
        if stripped.startswith("- "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.25
            apply_numbering(paragraph, bullet_num_id)
            add_inline(paragraph, stripped[2:].strip(), 11, INK)
            index += 1
            continue
        ordered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if ordered:
            if not in_ordered:
                ordered_num_id = add_numbering_definition(document, "decimal")
                in_ordered = True
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.25
            apply_numbering(paragraph, ordered_num_id or 1)
            add_inline(paragraph, ordered.group(1), 11, INK)
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                break
            if (
                candidate.startswith(("#", "- ", ">", "|", "```", "<!--"))
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        add_body_paragraph(" ".join(paragraph_lines).replace("  ", " "))


def audit_document(document: Document) -> None:
    section = document.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    assert all(round(value.inches, 2) == 1.0 for value in (section.top_margin, section.right_margin, section.bottom_margin, section.left_margin))
    required = {"Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"}
    assert required.issubset({style.name for style in document.styles})
    for table in document.tables:
        grid = table._tbl.tblGrid
        widths = [int(col.get(qn("w:w"))) for col in grid.findall(qn("w:gridCol"))]
        assert sum(widths) == CONTENT_WIDTH_DXA, widths
        tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
        assert tbl_w is not None and int(tbl_w.get(qn("w:w"))) == CONTENT_WIDTH_DXA
        tbl_ind = table._tbl.tblPr.find(qn("w:tblInd"))
        assert tbl_ind is not None and int(tbl_ind.get(qn("w:w"))) == TABLE_INDENT_DXA


def build() -> Path:
    figures = generate_figures()
    source_text = SOURCE.read_text(encoding="utf-8")
    document = Document()
    setup_styles(document)
    set_page_geometry(document)
    configure_header_footer(document.sections[0])
    document.core_properties.title = "GraphRAG 与 SAADS-Red-team 项目技术说明"
    document.core_properties.subject = "面向大模型应用安全防御的可审计 GraphRAG 工程"
    document.core_properties.author = "SAADS-Red-team"
    document.core_properties.keywords = "GraphRAG, LLM Security, Red Team, Knowledge Graph, RAG"
    add_cover(document)
    parse_markdown(document, source_text, figures)
    audit_document(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
